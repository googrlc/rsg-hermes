"""Tests for commission statement upload, staging and commit.

Two guards carry this feature and both get hammered here: the content-hash
dedupe (a statement must not be committable twice) and the crosscheck (a parse
that disagrees with the carrier's own stated total is wrong, and approving it
would put fiction in the ledger).
"""

from __future__ import annotations

from decimal import Decimal

import pytest

from hermes.commissions import statements as st

# A real Progressive statement's column set, taken from the raw_row of a line
# already in commission_transactions.
PROGRESSIVE_CSV = b"""policy_number,insured_name,prod,tran_code,tran_date,month_end,gross_premium,comm_rate,gross_comm,net_due_agent,producer,agent_code
862392084,HARRIS C.,Auto,New Business,08/29/2025,202508,1400,0.15,210,210,"COATES, GRETCHEN",03RC9
862392084,HARRIS C.,Auto,Renewal,02/28/2026,202602,1453,0.1,145.30,145.30,"COATES, GRETCHEN",03RC9
864561433,BROWN JOHNSO S.,Auto,Credit Endorsement,03/20/2026,202603,-7.08,0.1,-0.71,-0.71,"COATES, GRETCHEN",03RC9
"""


# --- money / rate / date parsing ---------------------------------------------

@pytest.mark.parametrize("raw,expected", [
    ("210", Decimal("210")), ("$1,453.00", Decimal("1453.00")),
    ("-0.71", Decimal("-0.71")), ("(75.50)", Decimal("-75.50")),
    ("", None), (None, None), ("-", None), ("junk", None),
])
def test_as_money(raw, expected):
    assert st.as_money(raw) == expected


def test_parenthesised_negatives_are_credits_not_positives():
    """A carrier writing (75.50) means minus. Reading it as +75.50 turns a
    chargeback into income."""
    assert st.as_money("(75.50)") == Decimal("-75.50")


@pytest.mark.parametrize("raw,expected", [
    ("0.15", Decimal("0.15")), ("15%", Decimal("0.15")), ("15", Decimal("0.15")),
    ("0.1", Decimal("0.1")), ("", None), ("junk", None),
])
def test_as_rate_normalizes_to_a_fraction(raw, expected):
    assert st.as_rate(raw) == expected


def test_a_bare_number_above_one_is_a_percentage():
    """No carrier pays 1500%. Reading 15 as 15.0 would corrupt every comparison."""
    assert st.as_rate("15") == Decimal("0.15")


@pytest.mark.parametrize("raw", ["2026-02-28", "02/28/2026", "20260228", "02-28-2026"])
def test_as_date_accepts_the_formats_carriers_use(raw):
    assert st.as_date(raw).isoformat() == "2026-02-28"


def test_as_date_returns_none_rather_than_guessing():
    assert st.as_date("last tuesday") is None


def test_month_key_falls_back_to_the_transaction_date():
    from datetime import date
    assert st.as_month_key("", date(2026, 2, 28)) == 202602
    assert st.as_month_key("202602") == 202602
    assert st.as_month_key("2026-02") == 202602
    assert st.as_month_key("", None) is None


@pytest.mark.parametrize("code,kind", [
    ("New Business", "new"), ("Renewal", "renewal"),
    # A credit endorsement is a premium credit, not a policy endorsement — the
    # more specific rule has to win over the bare "endorsement" one.
    ("Credit Endorsement", "adjustment"), ("Endorsement", "endorsement"),
    ("Chargeback", "chargeback"), ("Cancellation", "cancel"),
    ("Cancel Pro Rate", "cancel"), ("Reinstatement", "reinstatement"),
    # "new" is a substring of "renewal"; the renewal rule must be reached first.
    ("Renewal Business", "renewal"), ("New", "new"),
    ("Something Else", "other"), ("", "other"),
])
def test_transaction_type_normalization(code, kind):
    """The vocabulary matches what commission_transactions already holds.

    This path is the only writer now, so a value it invents here would show up
    beside 182 rows that name the same event differently.
    """
    assert st.normalize_transaction_type(code) == kind


# --- column aliasing ---------------------------------------------------------

def test_aliases_survive_carrier_column_naming():
    row = {"Policy #": "P1", "Gross Comm": "100", "Tran Date": "01/02/2026"}
    parsed = st.parse_row(row)
    assert parsed["policy_number"] == "P1"
    assert parsed["commission_amount"] == Decimal("100")
    assert parsed["transaction_date"] == "2026-01-02"


def test_a_row_with_neither_policy_nor_amount_is_not_a_line():
    assert st.parse_row({"producer": "COATES, GRETCHEN", "notes": "subtotal"}) is None


def test_the_original_row_is_kept_as_evidence():
    row = {"policy_number": "P1", "gross_comm": "100", "agent_code": "03RC9"}
    assert st.parse_row(row)["raw_row"]["agent_code"] == "03RC9"


# --- parsing a real statement ------------------------------------------------

def test_parses_a_progressive_statement():
    lines, warnings = st.parse_statement(PROGRESSIVE_CSV, "progressive-202602.csv")
    assert len(lines) == 3
    assert not [w for w in warnings if "skipped" not in w]

    new, renewal, credit = lines
    assert (new["transaction_type"], new["commission_amount"]) == ("new", Decimal("210"))
    assert (renewal["transaction_type"], renewal["month_key"]) == ("renewal", 202602)
    assert credit["transaction_type"] == "adjustment"
    assert credit["commission_amount"] == Decimal("-0.71")
    assert credit["lob"] == "Auto"          # from the "prod" column


def test_unsupported_formats_say_so_rather_than_returning_nothing():
    for name in ("x.xlsx", "x.pdf", "x.docx"):
        lines, warnings = st.parse_statement(b"data", name)
        assert lines == [] and warnings


# --- crosscheck --------------------------------------------------------------

def test_crosscheck_totals_the_parse():
    lines, _ = st.parse_statement(PROGRESSIVE_CSV, "s.csv")
    check = st.crosscheck(lines, stated_commission="354.59")
    assert check.parsed_commission == Decimal("354.59")
    assert check.ok and check.verifiable


def test_crosscheck_fails_when_the_parse_disagrees_with_the_carrier():
    lines, _ = st.parse_statement(PROGRESSIVE_CSV, "s.csv")
    check = st.crosscheck(lines, stated_commission="500.00")
    assert not check.ok
    assert check.commission_delta == Decimal("-145.41")


def test_no_stated_total_means_nothing_to_verify_against():
    lines, _ = st.parse_statement(PROGRESSIVE_CSV, "s.csv")
    check = st.crosscheck(lines)
    assert check.ok and not check.verifiable    # ok, but unproven


def test_a_cent_of_rounding_is_tolerated():
    lines, _ = st.parse_statement(PROGRESSIVE_CSV, "s.csv")
    assert st.crosscheck(lines, stated_commission="354.60").ok


# --- staging -----------------------------------------------------------------

class FakeSupa:
    def __init__(self, batches=None, ledger=None):
        self.tables = {
            st.BATCHES_TABLE: batches or [],
            st.STAGING_TABLE: [],
            st.STATEMENTS_TABLE: [],
            st.TRANSACTIONS_TABLE: [],
            "commission_ledger": ledger or [],
        }
        self._seq = 0
        self.updates: list[tuple[str, str, dict]] = []

    def select(self, table, *, columns="*", params=None, limit=1000):
        rows = self.tables.get(table, [])
        for k, v in (params or {}).items():
            if k == "order":
                continue
            if isinstance(v, str) and v.startswith("eq."):
                rows = [r for r in rows if str(r.get(k)) == v[3:]]
        return [dict(r) for r in rows][:limit]

    def insert(self, table, payload):
        self._seq += 1
        row = {"id": f"{table[:3]}-{self._seq}", **payload}
        self.tables.setdefault(table, []).append(row)
        return dict(row)

    def update(self, table, record_id, payload):
        self.updates.append((table, record_id, payload))
        for row in self.tables.get(table, []):
            if row.get("id") == record_id:
                row.update(payload)
                return dict(row)
        return {"id": record_id, **payload}


@pytest.fixture
def supa(monkeypatch):
    monkeypatch.setattr("hermes_core.book.select_policies", lambda *a, **k: [])
    return FakeSupa()


def test_staging_writes_a_batch_and_its_lines_but_no_money(supa):
    staged = st.stage_statement(
        supa, content=PROGRESSIVE_CSV, filename="p.csv",
        uploaded_by="lamar@risksolutionsgroup.net", carrier="Progressive",
        stated_commission="354.59",
    )
    assert staged.status == st.STATUS_PENDING_REVIEW
    assert staged.line_count == 3 and staged.approvable
    assert len(supa.tables[st.STAGING_TABLE]) == 3
    assert supa.tables[st.TRANSACTIONS_TABLE] == []    # nothing committed
    assert supa.tables[st.STATEMENTS_TABLE] == []


def test_reupload_of_the_same_file_is_refused_by_the_hash(supa):
    kw = dict(content=PROGRESSIVE_CSV, filename="p.csv",
              uploaded_by="lamar@risksolutionsgroup.net")
    first = st.stage_statement(supa, **kw)
    second = st.stage_statement(supa, **kw)
    assert second.duplicate_of == first.batch_id
    assert not second.approvable
    assert len(supa.tables[st.BATCHES_TABLE]) == 1     # no second batch
    assert len(supa.tables[st.STAGING_TABLE]) == 3     # no second set of lines


def test_a_failed_crosscheck_blocks_approval(supa):
    staged = st.stage_statement(
        supa, content=PROGRESSIVE_CSV, filename="p.csv",
        uploaded_by="lamar@risksolutionsgroup.net", stated_commission="9999",
    )
    assert not staged.crosscheck.ok
    assert not staged.approvable
    assert any("do not approve" in w for w in staged.warnings)


def test_an_unparseable_file_is_staged_as_error_not_silently_empty(supa):
    staged = st.stage_statement(
        supa, content=b"not,a,statement\n", filename="junk.csv",
        uploaded_by="lamar@risksolutionsgroup.net",
    )
    assert staged.status == st.STATUS_ERROR
    assert not staged.approvable
    assert any("no statement lines" in w for w in staged.warnings)


def test_the_preview_says_where_lines_would_land(supa):
    supa.tables["commission_ledger"] = [{"id": "L1", "policy_number": "862392084"}]
    staged = st.stage_statement(
        supa, content=PROGRESSIVE_CSV, filename="p.csv",
        uploaded_by="lamar@risksolutionsgroup.net",
    )
    assert staged.preview["will_link"] == 2          # both 862392084 lines
    assert staged.preview["will_be_unmatched"] == 1  # 864561433
    assert staged.preview["negative_lines"] == 1


# --- commit ------------------------------------------------------------------

def _stage(supa, **kw):
    return st.stage_statement(
        supa, content=PROGRESSIVE_CSV, filename="p.csv",
        uploaded_by="lamar@risksolutionsgroup.net", **kw)


def test_commit_refuses_a_batch_that_failed_its_crosscheck(supa):
    staged = _stage(supa, stated_commission="9999")
    with pytest.raises(ValueError, match="crosscheck"):
        st.commit_statement(supa, batch_id=staged.batch_id,
                            approved_by="lamar@risksolutionsgroup.net")


def test_commit_refuses_an_unknown_batch(supa):
    with pytest.raises(ValueError, match="not found"):
        st.commit_statement(supa, batch_id="nope", approved_by="a@b.net")


def test_commit_refuses_a_batch_twice(supa, monkeypatch):
    monkeypatch.setattr("hermes.commissions.matching.relink_unmatched",
                        lambda *a, **k: _NullLink())
    monkeypatch.setattr("hermes.commissions.reconcile.run_rollup",
                        lambda *a, **k: _NullRoll())
    staged = _stage(supa)
    st.commit_statement(supa, batch_id=staged.batch_id, approved_by="a@b.net")
    with pytest.raises(ValueError, match="committed"):
        st.commit_statement(supa, batch_id=staged.batch_id, approved_by="a@b.net")


class _NullLink:
    exact = normalized = ledger_rows_created = unmatched = 0
    errors: list = []


class _NullRoll:
    message = "rollup: noop"


def test_commit_promotes_lines_and_records_the_approver(supa, monkeypatch):
    monkeypatch.setattr("hermes.commissions.matching.relink_unmatched",
                        lambda *a, **k: _NullLink())
    monkeypatch.setattr("hermes.commissions.reconcile.run_rollup",
                        lambda *a, **k: _NullRoll())
    staged = _stage(supa, stated_commission="354.59")
    out = st.commit_statement(supa, batch_id=staged.batch_id,
                              approved_by="lamar@risksolutionsgroup.net")

    assert out.committed == 3
    assert len(supa.tables[st.TRANSACTIONS_TABLE]) == 3
    assert len(supa.tables[st.STATEMENTS_TABLE]) == 1

    batch = supa.tables[st.BATCHES_TABLE][0]
    assert batch["ingest_status"] == st.STATUS_COMMITTED
    assert batch["reviewed_by"] == "lamar@risksolutionsgroup.net"
    assert batch["statement_id"] == out.statement_id


def test_the_credit_line_commits_with_its_negative_amount(supa, monkeypatch):
    """is_negative is a GENERATED column, so we send the AMOUNT and let Postgres
    derive the flag. The signed amount is what has to survive the commit."""
    monkeypatch.setattr("hermes.commissions.matching.relink_unmatched",
                        lambda *a, **k: _NullLink())
    monkeypatch.setattr("hermes.commissions.reconcile.run_rollup",
                        lambda *a, **k: _NullRoll())
    staged = _stage(supa)
    st.commit_statement(supa, batch_id=staged.batch_id, approved_by="a@b.net")
    credits = [r for r in supa.tables[st.TRANSACTIONS_TABLE]
               if (r.get("commission_amount") or 0) < 0]
    assert len(credits) == 1
    assert credits[0]["policy_number"] == "864561433"


# --- reject ------------------------------------------------------------------

def test_reject_records_who_and_why(supa):
    staged = _stage(supa)
    st.reject_statement(supa, batch_id=staged.batch_id,
                        reviewed_by="lamar@risksolutionsgroup.net", reason="wrong month")
    batch = supa.tables[st.BATCHES_TABLE][0]
    assert batch["ingest_status"] == st.STATUS_REJECTED
    assert batch["flags"]["rejected_reason"] == "wrong month"


def test_a_rejected_batch_keeps_its_staged_lines_for_diagnosis(supa):
    staged = _stage(supa)
    st.reject_statement(supa, batch_id=staged.batch_id, reviewed_by="a@b.net")
    assert len(supa.tables[st.STAGING_TABLE]) == 3


def test_lines_inherit_the_statements_carrier(supa):
    """carrier_name is NOT NULL on staging and on commission_transactions, but
    statement lines rarely repeat it — it's a property of the statement. Caught
    live 2026-07-26 by a 23502 not-null violation."""
    st.stage_statement(supa, content=PROGRESSIVE_CSV, filename="p.csv",
                       uploaded_by="a@b.net", carrier="Progressive")
    assert all(r["carrier_name"] == "Progressive"
               for r in supa.tables[st.STAGING_TABLE])


def test_lines_never_stage_a_null_carrier_even_with_none_supplied(supa):
    st.stage_statement(supa, content=PROGRESSIVE_CSV, filename="p.csv",
                       uploaded_by="a@b.net", carrier=None)
    assert all(r["carrier_name"] for r in supa.tables[st.STAGING_TABLE])


# --- generated columns -------------------------------------------------------

def test_generated_columns_are_stripped_before_insert():
    """Postgres computes these; sending one raises 428C9. Bitten twice live —
    commission_ledger.delta and commission_transactions.is_negative."""
    payload = {"policy_number": "P1", "commission_amount": -5, "is_negative": True}
    out = st.strip_generated(st.TRANSACTIONS_TABLE, payload)
    assert "is_negative" not in out
    assert out["commission_amount"] == -5


def test_strip_generated_leaves_other_tables_alone():
    assert st.strip_generated("some_other_table", {"is_negative": True}) == {"is_negative": True}


def test_ledger_generated_columns_are_known():
    assert st.GENERATED_COLUMNS["commission_ledger"] == frozenset({"delta", "unearned_balance"})


def test_committed_lines_never_carry_a_generated_column(supa, monkeypatch):
    monkeypatch.setattr("hermes.commissions.matching.relink_unmatched",
                        lambda *a, **k: _NullLink())
    monkeypatch.setattr("hermes.commissions.reconcile.run_rollup",
                        lambda *a, **k: _NullRoll())
    staged = _stage(supa)
    st.commit_statement(supa, batch_id=staged.batch_id, approved_by="a@b.net")
    assert all("is_negative" not in r for r in supa.tables[st.TRANSACTIONS_TABLE])


def test_a_batch_where_every_line_failed_is_not_marked_committed(supa, monkeypatch):
    """Marking it committed would hide the failure and block a retry."""
    monkeypatch.setattr("hermes.commissions.matching.relink_unmatched",
                        lambda *a, **k: _NullLink())
    monkeypatch.setattr("hermes.commissions.reconcile.run_rollup",
                        lambda *a, **k: _NullRoll())
    staged = _stage(supa)

    def boom(table, payload):
        if table == st.TRANSACTIONS_TABLE:
            raise RuntimeError("insert rejected")
        return FakeSupa.insert(supa, table, payload)

    supa.insert = boom
    out = st.commit_statement(supa, batch_id=staged.batch_id, approved_by="a@b.net")
    assert out.committed == 0 and len(out.errors) == 3
    assert supa.tables[st.BATCHES_TABLE][0]["ingest_status"] == st.STATUS_ERROR


# --- xlsx + totals rows ------------------------------------------------------

def _workbook(rows, title="Statement", extra_sheet=False):
    from openpyxl import Workbook
    import io as _io
    wb = Workbook(); ws = wb.active; ws.title = title
    for row in rows:
        ws.append(row)
    if extra_sheet:
        wb.create_sheet("Notes")
    buf = _io.BytesIO(); wb.save(buf)
    return buf.getvalue()


HEADER = ["Policy #", "Insured Name", "Prod", "Tran Code", "Tran Date",
          "Gross Premium", "Comm Rate", "Gross Comm"]


def test_xlsx_is_parsed():
    content = _workbook([HEADER,
                         ["862392084", "HARRIS C.", "Auto", "Renewal", "2026-02-28", 1453, 0.1, 145.30]])
    lines, warnings = st.parse_statement(content, "s.xlsx")
    assert len(lines) == 1
    assert lines[0]["commission_amount"] == Decimal("145.3")
    assert lines[0]["transaction_type"] == "renewal"


def test_xlsx_header_is_found_not_assumed():
    """Carriers put a title block above the header. Assuming row 1 would parse a
    banner as column names and yield nothing."""
    content = _workbook([["PROGRESSIVE COMMISSION STATEMENT"], [], HEADER,
                         ["862392084", "H.", "Auto", "Renewal", "2026-02-28", 1453, 0.1, 145.30]])
    lines, warnings = st.parse_statement(content, "s.xlsx")
    assert len(lines) == 1
    assert any("header found on row 3" in w for w in warnings)


def test_xlsx_with_no_recognisable_header_says_so():
    content = _workbook([["just", "some", "text"], ["more", "text", "here"]])
    lines, warnings = st.parse_statement(content, "s.xlsx")
    assert lines == []
    assert any("could not find a header row" in w for w in warnings)


def test_multiple_sheets_are_flagged():
    content = _workbook([HEADER, ["P1", "X", "Auto", "Renewal", "2026-01-01", 100, 0.1, 10]],
                        extra_sheet=True)
    _, warnings = st.parse_statement(content, "s.xlsx")
    assert any("only 'Statement' was parsed" in w for w in warnings)


def test_a_totals_row_is_excluded_not_counted_as_income():
    """THE bug this guards: a TOTAL row carries an amount but no policy number.
    Counting it doubled a 2-line statement from $144.59 to $289.18."""
    content = _workbook([HEADER,
                         ["862392084", "H.", "Auto", "Renewal", "2026-02-28", 1453, 0.1, 145.30],
                         ["864561433", "B.", "Auto", "Credit Endorsement", "2026-03-20", -7.08, 0.1, -0.71],
                         ["", "", "", "", "TOTAL", "", "", 144.59]])
    lines, warnings = st.parse_statement(content, "s.xlsx")
    assert len(lines) == 2
    assert st.crosscheck(lines, stated_commission="144.59").ok
    assert any("NO policy number" in w for w in warnings)


def test_a_totals_row_in_csv_is_also_excluded():
    csv_bytes = (b"policy_number,gross_comm\n"
                 b"P1,100.00\n"
                 b",100.00\n")
    lines, warnings = st.parse_statement(csv_bytes, "s.csv")
    assert len(lines) == 1
    assert any("NO policy number" in w for w in warnings)


def test_excluded_totals_rows_are_named_so_a_human_can_check():
    csv_bytes = b"policy_number,gross_comm\n,55.00\n"
    _, warnings = st.parse_statement(csv_bytes, "s.csv")
    assert any("55.00" in w for w in warnings)


def test_legacy_xls_is_refused_with_a_route_forward():
    _, xls = st.parse_statement(b"x", "s.xls")
    assert any("re-save as .xlsx" in w for w in xls)


def test_an_unopenable_pdf_says_so_rather_than_blaming_the_build():
    """A corrupt file and a missing library are different problems.

    One is re-exported by whoever holds the statement; the other is fixed on the
    box. Reporting either as the other sends them to the wrong place.
    """
    lines, warnings, method = st.parse_upload(b"not a pdf at all", "s.pdf")
    assert lines == []
    assert method == st.METHOD_PDF_TEXT
    assert any("could not be opened as a PDF" in w for w in warnings)
    assert not any("PyMuPDF is not installed" in w for w in warnings)
    assert any("export it as CSV" in w for w in warnings)
